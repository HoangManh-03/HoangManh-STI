#!/usr/bin/env python3
"""
Msg description
    - sensor_msgs/LaserScan
      + angle_min: -3.1415927410125732 rad
      + angle_max: 3.1415927410125732 rad 
      + angle_increment: 0.004363323096185923 rad
      + time_increment: 3.4722223062999547e-05 seconds
      + scan_time: 0.05000000074505806 seconds
      + range_min: 0.0 m
      + range_max: 30.0 m
      + lenght of ranges: 1440
      + lenght of intensities: 1440
"""

from sensor_msgs.msg import PointCloud2, LaserScan
from std_msgs.msg import Int8
 
from math import atan2, sin, cos, sqrt, fabs, degrees, isnan, radians
from math import pi as PI
import rospy
import time
import copy

import json

import numpy as np
from scipy.optimize import minimize

from visualization_msgs.msg import Marker
from geometry_msgs.msg import Point
from navigation_reflector.msg import Raw_reflector
from openpyxl import Workbook

'''
- Đọc và ghi dữ liệu data vào excel tại thời điểm t
    + no reflector
    + haS reflector

- Check intensity of reflector in following distance
'''

class Excel_node():
    def __init__(self):
        rospy.init_node('excel_node', anonymous = True)
        self.rate = rospy.Rate(50)
        
        # -- param
        self.radius_reflector = rospy.get_param('~radius_reflector', 0.03)
        self.tolerance_radiusRelector = 0.01

        self.diameter_reflector = 2.0*self.radius_reflector
        self.circumference_reflector = 2.0*self.radius_reflector*PI
        self.half_circumference_reflector = self.radius_reflector*PI

        self.min_scanningRadius = rospy.get_param('~min_scanningRadius', 0.1)
        self.max_scanningRadius = rospy.get_param('~max_scanningRadius', 30.)

        self.min_reflectionIntensity = rospy.get_param('~min_reflectionIntensity', 1200)  # 970
        self.max_reflectionIntensity = rospy.get_param('~max_reflectionIntensity', 5000)

        # -- 
        self.x_base_to_lidar = rospy.get_param('~x_base_to_lidar', 0.361)
        self.y_base_to_lidar = rospy.get_param('~y_base_to_lidar', 0.261)
        self.r_base_to_lidar = rospy.get_param('~r_base_to_lidar', radians(45.))

        # -------- Topic Sub -------- #

        rospy.Subscriber("/scan_lms100", LaserScan, self.callback_scan)
        self.dataScan = LaserScan()
        self.is_scan = False

        # -------- Topic Pub -------- #
        # -- 
        self.process = 0
        self.record_status = 0

        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Data scanning from R2000'

        # Ghi tiêu đề cột
        self.ws.cell(row=1, column=1, value="Ranges")
        self.ws.cell(row=1, column=2, value="Intensity")

    def callback_scan(self, data):
        self.dataScan = data

        # print(len(self.dataScan.intensities))
        self.is_scan = True

    def run(self):
        while not rospy.is_shutdown():
            if self.process == 0:
                if self.is_scan:
                    self.process = 1

            elif self.process == 1:
                if self.is_scan == False:
                    self.rate.sleep()
                    continue
                
                # reset 
                self.is_scan = False

                # data_scan = copy.copy(self.dataScan)

                if self.record_status == 0:
                    self.record_status = 1

                    print(len(self.dataScan.ranges))
                    print(len(self.dataScan.intensities))

                    # Ghi dữ liệu vào cột
                    for i, (r, inten) in enumerate(zip(self.dataScan.ranges, self.dataScan.intensities), start=2):
                        self.ws.cell(row=i, column=1, value=r)      # Cột Ranges
                        self.ws.cell(row=i, column=2, value=inten) # Cột Intensity
                    
                    # Lưu file Excel
                    file_name = "Data_with_reflector_relation_di3.xlsx"
                    self.wb.save(file_name)

                    print(f"Dữ liệu đã được ghi vào file '{file_name}'.")

                # -- 
                # dis_arr = list(data_scan.ranges)

                # # -- lọc các điểm có cường độ cao 
                # ls_highReflectionIntensity = []
                # for index, hri in enumerate(data_scan.intensities):
                #     condition_1 = data_scan.ranges[index] >= self.min_scanningRadius
                #     condition_2 = hri >= self.min_reflectionIntensity and hri <= self.max_reflectionIntensity
                    
                #     if condition_1 and condition_2:
                #         ls_highReflectionIntensity.append((index, data_scan.ranges[index], hri))

                #     else:
                #         dis_arr[index] = 9.

                # data_scan.ranges = tuple(dis_arr)

                # self.pub_dataFilter.publish(data_scan)
                # print(ls_highReflectionIntensity) # (index, dis, intensity)
                # print('---+---')

                # -- gửi dữ liệu các điểm gương quét đươc

            self.rate.sleep()

def main():
    print ("--- Run Detect Reflector ---")
    program = Excel_node()
    program.run()

    print ("--- close! ---")

if __name__ == '__main__':
    main()




